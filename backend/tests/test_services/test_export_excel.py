import io
import pandas as pd
import pytest
from openpyxl import load_workbook

from src.services.export.export_excel import ExcelExport


class TestExcelExport:

    @pytest.fixture
    def sample_dataframe(self) -> pd.DataFrame:
        """Fixture que fornece um DataFrame fictício de auditoria com vários tipos de dados."""
        return pd.DataFrame(
            [
                {
                    "access_key": "35231000000000000000550010000000011000000000",
                    "modelo": "55",
                    "valor_cents": 15050,
                    "data_emissao": "2026-10-20",
                    "status_auditoria": "MISSING_XML",
                    "qtd_itens": 5,
                },
                {
                    "access_key": "35231000000000000000550010000000012000000000",
                    "modelo": "65",
                    "valor_cents": 20000,
                    "data_emissao": "2026-10-21",
                    "status_auditoria": "OK",
                    "qtd_itens": 2,
                },
            ]
        )

    def test_export_dataframe_to_bytes_empty_df_raises_value_error(self):
        """Valida se tentar exportar um DataFrame vazio dispara ValueError."""
        exporter = ExcelExport()
        with pytest.raises(ValueError, match="O DataFrame de auditoria está vazio"):
            exporter.export_dataframe_to_bytes(pd.DataFrame())

        with pytest.raises(ValueError, match="O DataFrame de auditoria está vazio"):
            exporter.export_dataframe_to_bytes(None)

    def test_export_dataframe_to_bytes_structure_and_headers(self, sample_dataframe):
        """Inspeciona o arquivo gerado em memória (BytesIO) e valida títulos e cabeçalhos."""
        exporter = ExcelExport(enterprise_name="Empresa Teste LTDA")
        excel_buffer = exporter.export_dataframe_to_bytes(sample_dataframe)

        assert isinstance(excel_buffer, io.BytesIO)

        wb = load_workbook(excel_buffer)
        assert "Auditoria Fiscal" in wb.sheetnames

        ws = wb["Auditoria Fiscal"]

        assert ws["A1"].value == "TRM Análises — RELATÓRIO DE AUDITORIA FISCAL"
        assert ws["A2"].value == "EMPRESA AUDITADA: EMPRESA TESTE LTDA"

        headers = [ws.cell(row=3, column=col).value for col in range(1, 7)]
        assert headers == ["ACCESS KEY", "MODELO", "VALOR", "DATA EMISSAO", "STATUS AUDITORIA", "QTD ITENS"]

    def test_export_dataframe_data_formatting(self, sample_dataframe):
        """Valida a conversão e formatação dos valores nas células de dados (a partir da linha 4)."""
        exporter = ExcelExport()
        excel_buffer = exporter.export_dataframe_to_bytes(sample_dataframe)
        wb = load_workbook(excel_buffer)
        ws = wb["Auditoria Fiscal"]

        assert ws["A4"].value == "35231000000000000000550010000000011000000000"
        assert ws["B4"].value == "NF-e"
        assert ws["C4"].value == 150.5
        assert ws["C4"].number_format == "R$ #,##0.00"
        assert ws["D4"].value == "2026-10-20"
        assert ws["F4"].value == 5
        assert ws["F4"].number_format == "#,##0"

    def test_export_dataframe_to_file(self, sample_dataframe, tmp_path):
        """Valida a gravação do arquivo físico em disco no diretório temporário."""
        output_file = tmp_path / "relatorios" / "teste_auditoria.xlsx"
        exporter = ExcelExport(enterprise_name="Empresa Fictícia")

        exporter.export_dataframe_to_file(sample_dataframe, output_file)

        assert output_file.exists()
        assert output_file.stat().st_size > 0


    def test_null_values_formatting(self):
        """Garante que valores nulos sejam exportados como células vazias no Excel."""
        df_nulls = pd.DataFrame([
            {"modelo": None, "valor_val": None, "status": None}
        ])

        exporter = ExcelExport()
        excel_buffer = exporter.export_dataframe_to_bytes(df_nulls)

        wb = load_workbook(excel_buffer)
        ws = wb["Auditoria Fiscal"]

        assert ws["A4"].value is None
        assert ws["B4"].value is None
        assert ws["C4"].value is None
